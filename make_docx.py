"""Build a Word version of the test manuscript, with the same planted problems.

Models how a Word user actually writes: a formatted reference list typed or
inserted by Zotero/EndNote, not a .bib file.
"""
from docx import Document
from docx.shared import Pt

d = Document()

d.add_heading("A sequence-based oracle for scFv binding and developability", 0)
d.add_paragraph("A. Montano Romero, [co-authors]")

d.add_heading("Abstract", level=1)
d.add_paragraph(
    "We present an oracle that predicts binding and developability properties of "
    "single-chain variable fragments (scFvs) from sequence alone. Trained on a "
    "30-member library characterised by a NanoBiT split-luciferase assay, our work "
    "achieves useful ranking performance on held-out clonal lineages. We release all "
    "code and data."
)

d.add_heading("Introduction", level=1)
d.add_paragraph(
    "Deep generative methods have transformed protein design over the past five years "
    "(Dauparas et al., 2022; Watson et al., 2022). Antibody-specific language models "
    "have followed (Olsen et al., 2022; Shuai et al., 2021). Property prediction, "
    "however, remains the bottleneck in most campaigns (Makowski et al., 2022)."
)
d.add_paragraph(
    "Earlier work established that developability properties can be co-optimised with "
    "affinity (Harvey et al., 2022), though the trade-offs are library-specific."
)

d.add_heading("Methods", level=1)
d.add_heading("Library design and expression", level=2)
d.add_paragraph(
    "Constructs were expressed in Expi293 cells. Full construct maps are deposited "
    "under GenBank accession MZ457891."
)
d.add_heading("Binding assay", level=2)
d.add_paragraph(
    "Binding was measured with a NanoBiT split-luciferase reporter. Raw plate data are "
    "available in the repository."
)
d.add_heading("Oracle training", level=2)
d.add_paragraph(
    "Descriptors were computed per sequence and a random forest trained with "
    "lineage-disjoint cross-validation. Code is available at "
    "https://github.com/montano-lab/scfv-binding-oracle."
)

d.add_heading("Results", level=1)
d.add_paragraph(
    "Our work recovers known binders at a rate well above the lineage-only baseline. "
    "The full comparison appears in Table 1."
)

d.add_heading("Data availability", level=1)
d.add_paragraph(
    "All code, trained models and processed assay data are available at "
    "https://github.com/montano-lab/scfv-binding-oracle. Analysis notebooks depend on "
    "the pipeline tooling described in Koenig et al. (2017)."
)

d.add_heading("References", level=1)
refs = [
    "Dauparas, J., Anishchenko, I., Bennett, N., et al. (2022). Robust deep learning "
    "based protein sequence design using ProteinMPNN. bioRxiv.",
    "Watson, J. L., Juergens, D., Bennett, N. R., et al. (2022). Broadly applicable and "
    "accurate protein design by integrating structure prediction networks and diffusion "
    "generative models. bioRxiv.",
    "Olsen, T. H., Moal, I. H., & Deane, C. M. (2022). AbLang: an antibody language "
    "model for completing antibody sequences. bioRxiv.",
    "Shuai, R. W., Ruffolo, J. A., & Gray, J. J. (2021). Generative language modeling "
    "for antibody design. bioRxiv.",
    "Makowski, E. K., Wu, L., Desai, A. A., & Tessier, P. M. (2022). Optimization of "
    "therapeutic antibodies for reduced self-association and non-specific binding. "
    "Protein Engineering, Design and Selection.",
    "Harvey, E. P., Shin, J.-E., Skiba, M. A., et al. (2022). An in silico method to "
    "assess antibody fragment polyreactivity. Nature Communications, 13.",
]
for r in refs:
    d.add_paragraph(r).paragraph_format.space_after = Pt(6)

d.save("paper.docx")
print("wrote paper.docx")
